"""Conversor markdown -> docx ligero para las hojas de lean ai.
Uso: python md2docx.py <entrada.md> [<entrada2.md> ...]
Genera un .docx junto a cada .md."""
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor

INLINE = [
    (re.compile(r"\*\*(.+?)\*\*"), r"\1"),
    (re.compile(r"`(.+?)`"), r"\1"),
    (re.compile(r"\[(.+?)\]\((.+?)\)"), r"\1"),
]


def clean(text):
    for pattern, repl in INLINE:
        text = pattern.sub(repl, text)
    return text


def convert(md_path):
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x4D)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and set(
            lines[i + 1].strip()
        ) <= set("|-: "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(
                    [clean(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                )
                i += 1
            rows = [rows[0]] + rows[2:]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    if c < len(table.rows[r].cells):
                        table.rows[r].cells[c].text = cell
            doc.add_paragraph()
            continue

        if not stripped:
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if heading:
            doc.add_heading(clean(heading.group(2)), level=len(heading.group(1)))
        elif stripped.startswith("> "):
            p = doc.add_paragraph(clean(stripped[2:]))
            p.style = "Intense Quote"
        elif re.match(r"^[-*]\s+", stripped):
            doc.add_paragraph(clean(re.sub(r"^[-*]\s+", "", stripped)), style="List Bullet")
        elif re.match(r"^\d+\.\s+", stripped):
            doc.add_paragraph(clean(re.sub(r"^\d+\.\s+", "", stripped)), style="List Number")
        elif set(stripped) <= set("-_*") and len(stripped) >= 3:
            doc.add_paragraph()
        else:
            doc.add_paragraph(clean(stripped))
        i += 1

    out = Path(md_path).with_suffix(".docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        print(convert(arg))
